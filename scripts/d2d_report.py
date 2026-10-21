"""
Build a D2D report (same columns as the Salesforce "Admit-Enrolled" export) for
every verified user in the Forklift DB.

Usage:
    python scripts/d2d_report.py [--db ~/Desktop/forklift.db] [--output data/d2d_report.xlsx]

"Devil to Devil Claimed Date" is the user's verified_at. "Visa Type" and
"Visited Campus" are not exposed by the CRM API and are left blank.
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
import time
import types
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import requests
from dotenv import dotenv_values

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

# Stub the package so salesforce.py loads without the discord voice stack.
_pkg = types.ModuleType("asu_discord")
_pkg.__path__ = [str(PROJECT_ROOT / "asu_discord")]
sys.modules.setdefault("asu_discord", _pkg)
from asu_discord import salesforce as sf  # noqa: E402

TARGET_TERM = "2267"
sf._target_term_codes = lambda: {TARGET_TERM}

HEADERS = [
    "Contact ID", "Opportunity ID", "Term: Term Name", "Type", "Stage", "Full Name",
    "Arizona Resident", "Home State/Province", "Visa Type",
    "Country of Citizenship: Country Name", "Location", "Devil to Devil Claimed Date",
    "Visited Campus", "ASURITE",
]
FAILED: list[str] = []
SEASONS = {"1": "Spring", "4": "Summer", "7": "Fall"}


def term_name(code: str | None) -> str | None:
    """ASU term code -> Salesforce term name, e.g. 2267 -> '2026 Fall'."""
    if not code or len(code) != 4:
        return code
    return f"{2000 + int(code[1:3])} {SEASONS.get(code[3], code[3])}"


def fetch(asurite: str, base: str, auth) -> tuple[dict | None, dict | None]:
    headers = auth()  # re-read each call so the ~1h OAuth token gets refreshed
    resp = requests.get(f"{base}/v1/crm-contact/contact?asurite={asurite}", headers=headers, timeout=15)
    resp.raise_for_status()
    contacts = resp.json().get("contact") or []
    if not contacts:
        return None, None
    contact = contacts[0]
    opps: list[dict] = []
    for career in ("Undergraduate", "Graduate"):
        r = requests.get(
            f"{base}/v1/crm-opportunity/opportunity?contactId={contact['id']}&career={career}",
            headers=headers, timeout=15,
        )
        if r.status_code == 200 and isinstance(r.json(), list):
            opps.extend(r.json())
    opps.sort(key=sf._created_date)
    opp, _ = sf._select_opportunity(opps)
    return contact, opp


def build_row(user: tuple, base: str, auth) -> list:
    asurite, full_name, verified_at = user
    asurite = asurite.strip().removesuffix("@asu.edu")
    claimed = datetime.fromisoformat(verified_at) if verified_at else None
    contact = opp = None
    for attempt in range(5):  # the CRM API 500s under load; back off and retry
        try:
            contact, opp = fetch(asurite, base, auth)
            break
        except requests.RequestException as exc:
            if attempt == 4:
                FAILED.append(asurite)
                print(f"{asurite}: {exc}", file=sys.stderr)
            else:
                time.sleep(2 ** attempt)
    contact, opp = contact or {}, opp or {}
    state = contact.get("state")
    international = sf._parse_bool(opp.get("internationalStudent"))
    name = f"{contact.get('firstname') or ''} {contact.get('lastname') or ''}".strip() or full_name
    return [
        (contact.get("id") or "")[:15] or None,          # SF reports use 15-char IDs
        (opp.get("opportunityId") or "")[:15] or None,
        term_name(opp.get("termCode")),
        opp.get("type"),
        opp.get("stageName"),
        name,
        # ponytail: approximated from contact state; CRM API has no tuition-residency field
        int((state or "").strip().lower() in {"arizona", "az"} and not international) if contact else None,
        state,
        None,
        contact.get("citizenshipCountry"),
        opp.get("locationName") or opp.get("currentLocation"),
        claimed,
        None,
        asurite,
    ]


def claim_bucket(d: datetime | None) -> str:
    if d is None:
        return "(none)"
    return "<1/15/2025" if d < datetime(2025, 1, 15) else str(d.year)


def write_xlsx(rows: list[list], output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Fall 2026 Admit-Enrolled"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for cell in ws["L"][1:]:
        cell.number_format = "m/d/yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # D2D Totals: count of Stage by claimed-date bucket, like the SF pivot.
    counts: dict[tuple[str, str], int] = {}
    for row in rows:
        key = (row[4] or "No Salesforce match", claim_bucket(row[11]))
        counts[key] = counts.get(key, 0) + 1
    stages = sorted({s for s, _ in counts})
    buckets = sorted({b for _, b in counts}, key=lambda b: (not b.startswith("<"), b))
    tot = wb.create_sheet("D2D Totals")
    tot.append([])
    tot.append([])
    tot.append(["Count of Stage", "Column Labels"])
    tot.append([None, *buckets, "Grand Total"])
    tot.append([])
    tot.append(["Row Labels"])
    for s in stages:
        vals = [counts.get((s, b), 0) for b in buckets]
        tot.append([s, *vals, sum(vals)])
    col_totals = [sum(counts.get((s, b), 0) for s in stages) for b in buckets]
    tot.append(["Grand Total", *col_totals, sum(col_totals)])

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument("--db", default=str(Path.home() / "Desktop" / "forklift.db"))
    parser.add_argument("--output", default=str(PROJECT_ROOT / "data" / "d2d_report.xlsx"))
    parser.add_argument("--env-file", default=str(PROJECT_ROOT / ".env"))
    parser.add_argument("--prod", action="store_true", help="Use prod CRM creds instead of nonprod.")
    parser.add_argument("--workers", type=int, default=4)
    args = parser.parse_args()

    env = dotenv_values(args.env_file)
    if args.prod:
        cid, secret, base, scope = env["SALESFORCE_API_CLIENT_ID"], env["SALESFORCE_API_CLIENT_SECRET"], sf._DEFAULT_BASE, None
    else:
        cid, secret = env["SALESFORCE_NONPROD_CLIENT_ID"], env["SALESFORCE_NONPROD_CLIENT_SECRET"]
        base = env.get("SALESFORCE_NONPROD_BASE_URL") or sf._DEFAULT_NONPROD_BASE
        scope = env.get("SALESFORCE_NONPROD_SCOPE_BASE_URL")
    token_cache: dict = {}

    def auth() -> dict:
        return {"Authorization": f"Bearer {sf._get_bearer_token_for(cid, secret, base, token_cache, scope_base_url=scope)}"}

    with sqlite3.connect(args.db) as conn:
        users = conn.execute(
            "SELECT asurite_id, full_name, verified_at FROM users WHERE verified = 1 ORDER BY verified_at"
        ).fetchall()
    print(f"Fetching {len(users)} verified users from {base}...", file=sys.stderr)

    with ThreadPoolExecutor(args.workers) as pool:
        rows = []
        for i, row in enumerate(pool.map(lambda u: build_row(u, base, auth), users), 1):
            rows.append(row)
            if i % 250 == 0:
                print(f"  {i}/{len(users)} ({len(FAILED)} failed)", file=sys.stderr)

    write_xlsx(rows, Path(args.output))
    matched = sum(1 for r in rows if r[1])
    print(f"Wrote {len(rows)} rows ({matched} with an opportunity) to {args.output}", file=sys.stderr)
    if FAILED:
        print(f"{len(FAILED)} users failed after retries (rows left blank): {', '.join(FAILED)}", file=sys.stderr)


if __name__ == "__main__":
    main()
